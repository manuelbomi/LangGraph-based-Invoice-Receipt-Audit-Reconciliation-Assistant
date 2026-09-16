"""Generates the synthetic demo documents used throughout this repo.

Produces:
  sample-data/invoices/*.pdf   -- 4 invoices, 3 distinct layouts (one is a
                                   near-duplicate resubmission of another, to
                                   demonstrate duplicate-payment detection)
  sample-data/receipts/*.png   -- 3 thermal-receipt-style images
  sample-data/ledger/general_ledger.xlsx -- a small GL / expense register

Everything here is entirely fictitious (vendors, amounts, people) -- see
sample-data/README.md for why these are synthesized rather than sourced from
a real company or a real invoice template.

Run with (from repo root):
    python sample-data/scripts/generate_samples.py

Requires: reportlab, pillow, openpyxl (all in backend/requirements.txt; for
regenerating outside the backend venv, `pip install reportlab pillow
openpyxl`).
"""
from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    HRFlowable,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

HERE = os.path.dirname(os.path.abspath(__file__))
SAMPLE_DATA_DIR = os.path.dirname(HERE)
INVOICES_DIR = os.path.join(SAMPLE_DATA_DIR, "invoices")
RECEIPTS_DIR = os.path.join(SAMPLE_DATA_DIR, "receipts")
LEDGER_DIR = os.path.join(SAMPLE_DATA_DIR, "ledger")

styles = getSampleStyleSheet()


# ---------------------------------------------------------------------------
# Invoice layout A: letterhead-style, left header, boxed meta, bordered table
# ---------------------------------------------------------------------------
def build_invoice_layout_a(path: str, *, invoice_number: str, invoice_date: str) -> None:
    doc = SimpleDocTemplate(path, pagesize=letter, topMargin=0.6 * inch, bottomMargin=0.6 * inch)
    story = []

    header_style = ParagraphStyle("Header", parent=styles["Heading1"], fontSize=18, textColor=colors.HexColor("#1a3c6e"))
    sub_style = ParagraphStyle("Sub", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#555555"))

    header_table = Table(
        [
            [
                Paragraph("BRIGHT OFFICE SUPPLIES INC.", header_style),
                Paragraph(f"<b>INVOICE</b><br/># {invoice_number}<br/>Date: {invoice_date}", sub_style),
            ]
        ],
        colWidths=[4.2 * inch, 2.3 * inch],
    )
    header_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(header_table)
    story.append(Paragraph("482 Industrial Parkway, Suite 12, Columbus, OH 43004", sub_style))
    story.append(Paragraph("Phone: (614) 555-0148  |  billing@brightofficesupplies.example", sub_style))
    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", color=colors.HexColor("#1a3c6e"), thickness=1.4))
    story.append(Spacer(1, 12))

    bill_to = Table(
        [
            [Paragraph("<b>Bill To:</b>", sub_style), Paragraph("<b>Payment Terms:</b>", sub_style)],
            [
                Paragraph("Meridian Analytics LLC<br/>900 Founders Row, Floor 4<br/>Chicago, IL 60601", sub_style),
                Paragraph("Net 30<br/>Currency: USD", sub_style),
            ],
        ],
        colWidths=[3.5 * inch, 3 * inch],
    )
    story.append(bill_to)
    story.append(Spacer(1, 16))

    line_items = [
        ("Copy Paper, Case of 10 Reams", "5", "42.00", "210.00"),
        ("Ergonomic Office Chair", "2", "249.00", "498.00"),
        ("Wireless Keyboard & Mouse Combo", "3", "35.00", "105.00"),
        ("Toner Cartridge (Black, High Yield)", "4", "89.25", "357.00"),
    ]
    table_data = [["Description", "Qty", "Unit Price", "Amount"]] + [list(r) for r in line_items]
    items_table = Table(table_data, colWidths=[3.2 * inch, 0.7 * inch, 1.1 * inch, 1.1 * inch])
    items_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a3c6e")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9.5),
                ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#999999")),
                ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f5fa")]),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(items_table)
    story.append(Spacer(1, 10))

    totals = Table(
        [
            ["", "Subtotal:", "$1,170.00"],
            ["", "Tax (8.25%):", "$96.53"],
            ["", "Total Due:", "$1,266.53"],
        ],
        colWidths=[3.6 * inch, 1.3 * inch, 1.2 * inch],
    )
    totals.setStyle(
        TableStyle(
            [
                ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                ("FONTNAME", (1, 2), (-1, 2), "Helvetica-Bold"),
                ("LINEABOVE", (1, 2), (-1, 2), 0.8, colors.HexColor("#1a3c6e")),
                ("FONTSIZE", (0, 0), (-1, -1), 10),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story.append(totals)
    story.append(Spacer(1, 20))
    story.append(
        Paragraph(
            "Thank you for your business. Please remit payment within 30 days to the "
            "address above, referencing the invoice number.",
            sub_style,
        )
    )
    doc.build(story)


def build_invoice_layout_b(path: str) -> None:
    """Summit Cloud Hosting LLC -- centered header, minimalist unbordered table."""
    doc = SimpleDocTemplate(path, pagesize=letter, topMargin=0.7 * inch, bottomMargin=0.6 * inch)
    story = []

    center_header = ParagraphStyle("CenterHeader", parent=styles["Heading1"], alignment=TA_CENTER, fontSize=20, textColor=colors.HexColor("#0f766e"))
    center_sub = ParagraphStyle("CenterSub", parent=styles["Normal"], alignment=TA_CENTER, fontSize=9.5, textColor=colors.HexColor("#444444"))
    left_style = ParagraphStyle("Left", parent=styles["Normal"], fontSize=9.5, alignment=TA_LEFT)
    right_style = ParagraphStyle("Right", parent=styles["Normal"], fontSize=9.5, alignment=TA_RIGHT)

    story.append(Paragraph("SUMMIT CLOUD HOSTING LLC", center_header))
    story.append(Paragraph("cloud infrastructure &amp; managed services", center_sub))
    story.append(Spacer(1, 6))
    story.append(HRFlowable(width="100%", color=colors.HexColor("#0f766e"), thickness=1.2))
    story.append(Spacer(1, 14))

    meta = Table(
        [
            [
                Paragraph("<b>Bill To</b><br/>Meridian Analytics LLC<br/>900 Founders Row, Floor 4<br/>Chicago, IL 60601", left_style),
                Paragraph("<b>Invoice Details</b><br/>Invoice #: INV-2044<br/>Date: 2026-08-10<br/>Terms: Net 15<br/>Currency: USD", right_style),
            ]
        ],
        colWidths=[3.3 * inch, 3.3 * inch],
    )
    story.append(meta)
    story.append(Spacer(1, 18))

    header_row = ["Service", "Qty", "Unit Price", "Amount"]
    rows = [
        ("Cloud Compute - Production Cluster (Monthly)", "1", "$2,400.00", "$2,400.00"),
        ("Object Storage - 5TB Tier", "1", "$350.00", "$350.00"),
        ("Premium Support Plan", "1", "$500.00", "$500.00"),
    ]
    table_data = [header_row] + [list(r) for r in rows]
    items_table = Table(table_data, colWidths=[3.4 * inch, 0.6 * inch, 1.1 * inch, 1.1 * inch])
    items_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("LINEBELOW", (0, 0), (-1, 0), 1, colors.HexColor("#0f766e")),
                ("LINEBELOW", (0, 1), (-1, -1), 0.4, colors.HexColor("#cccccc")),
                ("FONTSIZE", (0, 0), (-1, -1), 9.5),
                ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.append(items_table)
    story.append(Spacer(1, 12))

    totals_box = Table(
        [["Subtotal", "$3,250.00"], ["Tax (8.25%)", "$268.13"], ["Total Due", "$3,518.13"]],
        colWidths=[1.4 * inch, 1.4 * inch],
    )
    totals_box.setStyle(
        TableStyle(
            [
                ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#0f766e")),
                ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")),
                ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                ("FONTNAME", (0, 2), (-1, 2), "Helvetica-Bold"),
                ("BACKGROUND", (0, 2), (-1, 2), colors.HexColor("#e6f4f2")),
                ("FONTSIZE", (0, 0), (-1, -1), 10),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    wrapper = Table([[None, totals_box]], colWidths=[3.4 * inch, 2.8 * inch])
    story.append(wrapper)
    story.append(Spacer(1, 22))
    story.append(
        Paragraph(
            "Questions about this invoice? Contact billing@summitcloudhosting.example. "
            "Amounts over $3,000 require manager approval per standard AP policy.",
            center_sub,
        )
    )
    doc.build(story)


def build_invoice_layout_c(path: str) -> None:
    """Redwood Facilities Maintenance -- compact, top-right meta, shaded rows."""
    doc = SimpleDocTemplate(path, pagesize=letter, topMargin=0.6 * inch, bottomMargin=0.6 * inch)
    story = []

    small = ParagraphStyle("Small", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#333333"))
    title = ParagraphStyle("Title", parent=styles["Heading2"], fontSize=14, textColor=colors.HexColor("#7c2d12"))
    right_meta = ParagraphStyle("RightMeta", parent=styles["Normal"], fontSize=9, alignment=TA_RIGHT)

    top = Table(
        [[Paragraph("Redwood Facilities Maintenance", title), Paragraph("Invoice: FM-778<br/>Date: 2026-08-12", right_meta)]],
        colWidths=[4 * inch, 2.5 * inch],
    )
    story.append(top)
    story.append(Paragraph("1120 Willow Creek Rd, Portland, OR 97214 | (503) 555-0199", small))
    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", color=colors.HexColor("#7c2d12"), thickness=0.8))
    story.append(Spacer(1, 10))
    story.append(Paragraph("<b>Remit To:</b> same address above. <b>Client:</b> Meridian Analytics LLC, Chicago, IL", small))
    story.append(Spacer(1, 14))

    rows = [
        ("HVAC Quarterly Service", "1", "$380.00", "$380.00"),
        ("Janitorial Supplies Restock", "1", "$145.00", "$145.00"),
        ("Emergency Lock Repair", "1", "$87.75", "$87.75"),
    ]
    table_data = [["Item", "Qty", "Rate", "Amount"]] + [list(r) for r in rows]
    items_table = Table(table_data, colWidths=[3.2 * inch, 0.6 * inch, 1.1 * inch, 1.1 * inch])
    items_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#7c2d12")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 9.5),
                ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#fdf3ee"), colors.white]),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(items_table)
    story.append(Spacer(1, 10))

    totals = Table(
        [["Subtotal:", "$612.75"], ["Tax:", "$0.00 (service exempt)"], ["Total Due:", "$612.75"]],
        colWidths=[1.3 * inch, 2 * inch],
    )
    totals.setStyle(
        TableStyle(
            [
                ("ALIGN", (0, 0), (-1, -1), "RIGHT"),
                ("FONTNAME", (0, 2), (-1, 2), "Helvetica-Bold"),
                ("LINEABOVE", (0, 2), (-1, 2), 0.8, colors.HexColor("#7c2d12")),
                ("FONTSIZE", (0, 0), (-1, -1), 10),
            ]
        )
    )
    wrapper = Table([[None, totals]], colWidths=[3.5 * inch, 3.3 * inch])
    story.append(wrapper)
    story.append(Spacer(1, 18))
    story.append(Paragraph("<i>Due on receipt. Late payments accrue 1.5% monthly interest.</i>", small))
    doc.build(story)


# ---------------------------------------------------------------------------
# Receipts (thermal-receipt-style PNGs, drawn with PIL)
# ---------------------------------------------------------------------------
def _mono_font(size: int) -> ImageFont.FreeTypeFont:
    candidates = [
        "C:/Windows/Fonts/consola.ttf",
        "C:/Windows/Fonts/cour.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf",
    ]
    for path in candidates:
        if os.path.exists(path):
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def draw_receipt(
    path: str,
    *,
    merchant: str,
    address: str,
    date: str,
    items: list[tuple[str, str]],
    subtotal: str,
    tax: str,
    total: str,
    payment_method: str,
    receipt_no: str,
) -> None:
    width = 380
    line_h = 22
    height = 260 + line_h * (len(items) + 6)
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font = _mono_font(14)
    font_bold = _mono_font(16)

    y = 14

    def center(text: str, f=font):
        bbox = draw.textbbox((0, 0), text, font=f)
        w = bbox[2] - bbox[0]
        draw.text(((width - w) / 2, y), text, font=f, fill="black")

    center(merchant, font_bold)
    y += line_h
    center(address)
    y += line_h
    center(f"Receipt #{receipt_no}")
    y += line_h
    center(date)
    y += line_h + 6
    draw.text((14, y), "-" * 44, font=font, fill="black")
    y += line_h

    for name, amount in items:
        draw.text((14, y), name, font=font, fill="black")
        bbox = draw.textbbox((0, 0), amount, font=font)
        w = bbox[2] - bbox[0]
        draw.text((width - 14 - w, y), amount, font=font, fill="black")
        y += line_h

    draw.text((14, y), "-" * 44, font=font, fill="black")
    y += line_h

    for label, amount in [("Subtotal", subtotal), ("Tax", tax)]:
        draw.text((14, y), label, font=font, fill="black")
        bbox = draw.textbbox((0, 0), amount, font=font)
        w = bbox[2] - bbox[0]
        draw.text((width - 14 - w, y), amount, font=font, fill="black")
        y += line_h

    draw.text((14, y), "TOTAL", font=font_bold, fill="black")
    bbox = draw.textbbox((0, 0), total, font=font_bold)
    w = bbox[2] - bbox[0]
    draw.text((width - 14 - w, y), total, font=font_bold, fill="black")
    y += line_h + 6

    draw.text((14, y), "-" * 44, font=font, fill="black")
    y += line_h
    center(f"Paid via {payment_method}")
    y += line_h
    center("Thank you for your visit!")

    img.save(path)


# ---------------------------------------------------------------------------
# General ledger (.xlsx)
# ---------------------------------------------------------------------------
def build_ledger(path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "General Ledger"

    headers = ["date", "vendor", "description", "gl_account", "amount", "currency"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1A3C6E", end_color="1A3C6E", fill_type="solid")

    rows = [
        ("2026-08-05", "Bright Office Supplies Inc.", "Office supplies restock", "6100-Office Supplies", 1266.53, "USD"),
        ("2026-08-07", "Bright Office Supplies Inc.", "Office supplies restock (resubmit)", "6100-Office Supplies", 1266.53, "USD"),
        ("2026-08-10", "Summit Cloud Hosting LLC", "Monthly cloud infrastructure", "6400-IT Services", 3518.13, "USD"),
        ("2026-08-11", "Metro Taxi Co", "Ground transportation - client visit", "6550-Travel", 47.00, "USD"),
        ("2026-08-03", "Acme Legal Services", "Contract review retainer", "6700-Professional Fees", 2200.00, "USD"),
        ("2026-08-09", "Statewide Insurance Co", "General liability premium", "6800-Insurance", 875.00, "USD"),
    ]
    for row in rows:
        ws.append(row)

    for i, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(18, len(header) + 4)

    wb.save(path)


def main() -> None:
    os.makedirs(INVOICES_DIR, exist_ok=True)
    os.makedirs(RECEIPTS_DIR, exist_ok=True)
    os.makedirs(LEDGER_DIR, exist_ok=True)

    build_invoice_layout_a(
        os.path.join(INVOICES_DIR, "INV-1001_bright_office_supplies.pdf"),
        invoice_number="INV-1001",
        invoice_date="2026-08-05",
    )
    # Near-duplicate resubmission of INV-1001 -- same vendor, same amount,
    # date two days later, different invoice number. This is the intentional
    # duplicate-payment scenario the reconcile node must catch.
    build_invoice_layout_a(
        os.path.join(INVOICES_DIR, "INV-1006_bright_office_supplies_DUPLICATE.pdf"),
        invoice_number="INV-1006",
        invoice_date="2026-08-07",
    )
    build_invoice_layout_b(os.path.join(INVOICES_DIR, "INV-2044_summit_cloud_hosting.pdf"))
    build_invoice_layout_c(os.path.join(INVOICES_DIR, "FM-778_redwood_facilities.pdf"))

    draw_receipt(
        os.path.join(RECEIPTS_DIR, "downtown_cafe_receipt.png"),
        merchant="DOWNTOWN CAFE",
        address="214 Main St, Chicago, IL",
        date="2026-08-06 08:14",
        items=[("Latte", "$4.75"), ("Blueberry Muffin", "$3.25"), ("Avocado Toast", "$8.50")],
        subtotal="$16.50",
        tax="$1.36",
        total="$17.86",
        payment_method="VISA ****1234",
        receipt_no="48213",
    )
    draw_receipt(
        os.path.join(RECEIPTS_DIR, "quickmart_receipt.png"),
        merchant="QUICKMART CONVENIENCE",
        address="77 Harrison Ave, Chicago, IL",
        date="2026-08-09 17:42",
        items=[("Bottled Water", "$1.99"), ("Snack Mix", "$3.49"), ("Phone Charger Cable", "$12.99")],
        subtotal="$18.47",
        tax="$1.52",
        total="$19.99",
        payment_method="MASTERCARD ****5566",
        receipt_no="90871",
    )
    draw_receipt(
        os.path.join(RECEIPTS_DIR, "metro_taxi_receipt.png"),
        merchant="METRO TAXI CO",
        address="Chicago, IL - Ride Receipt",
        date="2026-08-11 19:05",
        items=[("Airport Transfer Fare", "$42.00"), ("Toll Fee", "$5.00")],
        subtotal="$47.00",
        tax="$0.00",
        total="$47.00",
        payment_method="CORPORATE CARD ****9021",
        receipt_no="TX-55210",
    )

    build_ledger(os.path.join(LEDGER_DIR, "general_ledger.xlsx"))

    print("Sample data generated in:", SAMPLE_DATA_DIR)


if __name__ == "__main__":
    main()
